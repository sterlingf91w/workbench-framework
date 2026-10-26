"""Single-user localhost workspace. The xlsx workbook is the source of truth."""
from copy import copy
from io import BytesIO
import base64
import re
import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import secrets
import shutil
import tempfile
import threading
from datetime import date, datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, unquote
import webbrowser

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import TableColumn, Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties
from docx import Document
from business import business_rows, refresh_business_sheet, STEP_ACTIONS

ROOT = Path(__file__).resolve().parent
WORK_CALENDAR = json.loads((ROOT / 'work-calendar.json').read_text())

def third_workday(received):
    day = date.fromisoformat(received)
    count = 0
    while count < 3:
        day += timedelta(days=1)
        if day.year not in WORK_CALENDAR['years']:
            return ''
        key = day.isoformat()
        if key in WORK_CALENDAR['workdays'] or (day.weekday() < 5 and key not in WORK_CALENDAR['holidays']):
            count += 1
    return day.isoformat()

CAPACITY = 2000
LAST_ROW = CAPACITY + 7
EXTRA_FIELDS = {
    '沟通记录': {'客户整理稿及时间戳': [], '我方整理稿及时间戳': [], '整理稿路径': []},
    '样品测试': {'延期原因': [], '推进记录': [], '申请时间': [], '图号': [], '物料名称': [], '计划测试日期': [], '规格': [], '业务员': [], '备注': []},
    '测试轮次': {'任务编号': [], '日期依据': ['客户约定', '本人计划', '待确认'], '延期原因': [], '推进记录': []},
    '用胶项目': {'商务进度': ['待确认', '采购路径确认中', '报价中', '条件确认中', '订单已确认'], '当前卡点': [], '等待对象': [],
        **{f: [] for f in ['开始日期','关闭日期','应用端客户名称','项目发起人','业务承接人','项目开放原因','竞品主要问题','包装规格','当前售价','可接受价格预估','项目外部现状','方案性质','项目备注','最新反馈','反馈日期','跟进历史','需要得到的支持']}},
    '跟进待办': {'延期原因': [], '日期调整记录': [], '动作来源': ['手动','样品自动','测试自动'], '自动步骤': []},
}
TEMPLATE = ROOT / 'template' / '空白跟进表.xlsx'
SHEETS = {'跟进待办': 'T', '公司': 'C', '联系人': 'P', '用胶项目': 'J',
          '沟通记录': 'V', '样品测试': 'S', '技术经验': 'K', '每日复盘': 'D', '测试轮次': 'R'}
TRIAL_FIELDS = ['测试编号', '项目编号', '样品编号', '日期', '测试负责人编号', '测试条件', '合格标准', '状态', '测试结果与原因', '证据路径或链接', '下一步具体动作', '约定反馈日期']
RELATIONS = {
    '测试轮次': {'项目编号': '用胶项目', '样品编号': '样品测试', '测试负责人编号': '联系人'},
    '联系人': {'公司编号': '公司'},
    '用胶项目': {'公司编号': '公司', '主要联系人编号': '联系人'},
    '跟进待办': {'公司编号': '公司', '联系人编号': '联系人', '项目编号': '用胶项目', '样品编号': '样品测试', '来源沟通编号': '沟通记录'},
    '沟通记录': {'公司编号': '公司', '联系人编号': '联系人', '项目编号': '用胶项目', '下一步任务编号': '跟进待办'},
    '样品测试': {'项目编号': '用胶项目', '收件联系人编号': '联系人', '测试负责人编号': '联系人', '任务编号': '跟进待办', '来源沟通编号': '沟通记录'},
    '技术经验': {'项目编号': '用胶项目', '来源沟通编号': '沟通记录'},
}
DATE_FIELDS = {'开始日期', '关闭日期', '反馈日期', '计划测试日期', '申请时间', '计划日期', '完成日期', '日期', '寄出日期', '签收日期', '约定反馈日期', '答复日期'}
NUMBER_FIELDS = {'数量', '客户生产用量'}
CHANNEL_LIST = '"面访,电话,微信,朋友圈,邮件,社交媒体,视频会议,其他"'
SAMPLE_STATUS = '"待确认,已申请,待寄出,已寄出,已签收,测试中,通过,不通过,暂停,取消"'
SHIPPED = ('已寄出', '已签收', '通过', '不通过')
TRANSCRIPT_MARKER = re.compile(r'^(\d{1,2}:\d{2})\s+([^：:]{1,30}?)\s+\[段落(\d+)\][：:](.*)$')
DOCX_SPEAKER = re.compile(r'^(.{1,30}?)\s+(\d{1,2}:\d{2})$')


def transcript_upload(body, record):
    if not isinstance(body.get('filename'), str) or not body['filename'].lower().endswith('.docx'):
        raise UserError('请上传已核对的 Word 文件（.docx）。')
    try:
        raw = base64.b64decode(body['content'], validate=True)
    except (KeyError, ValueError):
        raise UserError('文件无法读取，请重新选择。')
    if not raw or len(raw) > 1500000:
        raise UserError('校对稿需小于 1.5 MB。')
    try:
        paragraphs = [p.text.strip() for p in Document(BytesIO(raw)).paragraphs]
    except Exception:
        raise UserError('Word 文件无法解析，请确认它是有效的 .docx。')
    uploaded = []
    current = None
    for line in paragraphs:
        marker = DOCX_SPEAKER.fullmatch(line)
        if marker:
            if current:
                uploaded.append(current)
            current = {'speaker': marker[1].strip(), 'time': marker[2], 'text': ''}
        elif current and line:
            current['text'] += ('\n' if current['text'] else '') + line
    if current:
        uploaded.append(current)
    original = []
    for side, field in (('customer', '客户原话及时间戳'), ('ours', '我方原话及时间戳')):
        for line in str(record.get(field, '')).splitlines():
            marker = TRANSCRIPT_MARKER.match(line)
            if marker:
                original.append({'time': marker[1], 'speaker': marker[2].strip(), 'paragraph': int(marker[3]), 'side': side})
    original.sort(key=lambda row: row['paragraph'])
    if not original or len(uploaded) != len(original) or any(not u['text'] for u in uploaded):
        raise UserError(f'校对稿需保留原来的说话人、时间戳和 {len(original)} 段对话；请检查是否缺段或多段。')
    if [(u['speaker'], u['time']) for u in uploaded] != [(o['speaker'], o['time']) for o in original]:
        raise UserError('校对稿的说话人或时间戳顺序与原稿不一致。请沿用原稿标记，只修改对话正文。')
    for turn, source in zip(uploaded, original):
        turn.update(source)
    for side in ('customer', 'ours'):
        text = '\n'.join(f"{t['time']} {t['speaker']} [段落{t['paragraph']}]：{t['text']}" for t in uploaded if t['side'] == side)
        if len(text) > 30000:
            raise UserError('校对稿内容超出单侧 30000 字限制。')
    return raw, uploaded


class UserError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def serializable(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    return '' if v is None else v


def set_list(ws, sqref, formula):
    for dv in ws.data_validations.dataValidation:
        if sqref in str(dv.sqref).split():
            if dv.formula1 != formula:
                dv.formula1 = formula
                return True
    return False


class Store:
    def __init__(self, path, template=TEMPLATE):
        self.path = Path(path).resolve()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            shutil.copy2(template, self.path)
        self.backups = self.path.parent / 'backups'
        self.lock = threading.Lock()
        self.ensure_options()

    def ensure_options(self):
        version = digest(self.path)
        wb = load_workbook(self.path)
        changed = False
        if '测试轮次' not in wb.sheetnames:
            ws = wb.create_sheet('测试轮次')
            ws.cell(1, 1, '样品测试轮次：每次复测新增一条，保留历史结果')
            ws.freeze_panes = 'D8'
            for i, name in enumerate(TRIAL_FIELDS, 1):
                ws.cell(7, i, name)._style = copy(wb['样品测试'].cell(7, 1)._style)
                ws.cell(8, i)._style = copy(wb['样品测试'].cell(8, 2)._style)
                ws.column_dimensions[get_column_letter(i)].width = 28 if i > 5 else 20
            table = Table(displayName='TrialRounds', ref=f'A7:L{LAST_ROW}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(table)
            dv = DataValidation(type='list', formula1='"待安排,测试中,通过,不通过,待复测,取消"', allow_blank=True)
            dv.add(f'H8:H{LAST_ROW}')
            ws.add_data_validation(dv)
            changed = True
        for name in SHEETS:
            ws = wb[name]
            headers = [c.value for c in ws[7]]
            for field, choices in EXTRA_FIELDS.get(name, {}).items():
                if field in headers:
                    continue
                col = len(headers) + 1
                headers.append(field)
                ws.cell(7, col, field)._style = copy(ws.cell(7, col - 1)._style)
                ws.column_dimensions[get_column_letter(col)].width = 28
                if choices:
                    dv = DataValidation(type='list', formula1='"' + ','.join(choices) + '"', allow_blank=True)
                    dv.add(f'{get_column_letter(col)}8:{get_column_letter(col)}{LAST_ROW}')
                    ws.add_data_validation(dv)
                for table in ws.tables.values():
                    table.tableColumns.append(TableColumn(id=col, name=field))
                changed = True
            for row in ws:
                for c in row:
                    if c.data_type == 'f' and '$207' in c.value:
                        c.value = c.value.replace('$207', f'${LAST_ROW}')
                        changed = True
            for dv in ws.data_validations.dataValidation:
                before = str(dv.sqref)
                after = re.sub(r'(?<=[A-Z])207\b', str(LAST_ROW), before)
                if before != after:
                    dv.sqref = after
                    changed = True
            for cf in list(ws.conditional_formatting):
                before = str(cf.sqref)
                after = re.sub(r'(?<=[A-Z])207\b', str(LAST_ROW), before)
                if before != after:
                    rules = list(ws.conditional_formatting[cf])
                    del ws.conditional_formatting[before]
                    for rule in rules:
                        ws.conditional_formatting.add(after, rule)
                    changed = True
            for table in ws.tables.values():
                ref = f'A7:{get_column_letter(len(headers))}{LAST_ROW}'
                if table.ref != ref:
                    table.ref = ref
                    if table.autoFilter:
                        table.autoFilter.ref = ref
                    changed = True
        changed = set_list(wb['沟通记录'], f'F8:F{LAST_ROW}', CHANNEL_LIST) or changed
        changed = set_list(wb['样品测试'], f'Q8:Q{LAST_ROW}', SAMPLE_STATUS) or changed
        for row in wb['使用说明']:
            for c in row:
                if c.data_type != 'f' and isinstance(c.value, str) and '200' in c.value:
                    new = re.sub(r'(?<!\d)200(?!\d)', str(CAPACITY), c.value)
                    if new != c.value:
                        c.value = new
                        changed = True
        if '业务跟进表' not in wb.sheetnames:
            changed = True
        if not changed:
            return
        if any((self.path.parent / n).exists() for n in ('~$' + self.path.name, '.~lock.' + self.path.name + '#')):
            raise UserError('请关闭 Excel 后重新启动，以升级工作簿。', 409)
        wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
        fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
        os.close(fd)
        try:
            refresh_business_sheet(wb)
            wb.save(temp)
            load_workbook(temp).close()
            if digest(self.path) != version:
                raise UserError('升级期间 Excel 已变化，请重新启动。', 409)
            self.backups.mkdir(exist_ok=True)
            shutil.copy2(self.path, self.backups / f'升级前-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx')
            os.replace(temp, self.path)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)

    def read(self):
        version = digest(self.path)
        wb = load_workbook(self.path)
        data, schema = {}, {}
        for name in SHEETS:
            s = wb[name]
            headers = [c.value for c in s[7]]
            fields = []
            for i, h in enumerate(headers, 1):
                choices = []
                for dv in s.data_validations.dataValidation:
                    if f'{get_column_letter(i)}8' in str(dv.sqref) and dv.type == 'list' and dv.formula1:
                        choices = dv.formula1.strip('"').split(',')
                fields.append({'name': h, 'computed': '〔自动〕' in str(h), 'options': choices,
                               'type': 'date' if h in DATE_FIELDS else 'number' if h in NUMBER_FIELDS else 'text',
                               'readonly': h in ('日期调整记录', '推进记录', '跟进历史', '动作来源', '自动步骤') or (name in ('样品测试', '测试轮次') and h == '任务编号'), 'relation': RELATIONS.get(name, {}).get(h)})
            schema[name] = fields
            data[name] = []
            if any(c.value is not None and c.data_type != 'f' for row in s.iter_rows(min_row=LAST_ROW + 1) for c in row):
                raise UserError(f'{name}超出工作台容量后仍有数据，请先扩展工作台容量，避免遗漏。')
            for row in range(8, min(s.max_row, LAST_ROW) + 1):
                if not s.cell(row, 1).value:
                    if any(s.cell(row, i).value is not None and s.cell(row, i).data_type != 'f'
                           for i in range(2, len(headers) + 1)):
                        raise UserError(f'{name}第{row}行有内容但缺少编号，请先在Excel补齐。')
                    continue
                record = {h: serializable(s.cell(row, i).value) for i, h in enumerate(headers, 1)
                          if h and '〔自动〕' not in str(h)}
                for field in fields:
                    if not field['computed'] and field['type'] == 'text':
                        record[field['name']] = str(record[field['name']])
                record['_id'] = str(s.cell(row, 1).value)
                record['_row'] = row
                data[name].append(record)
        if digest(self.path) != version:
            raise UserError('Excel正在变化，请稍后刷新。', 409)
        return wb, {'version': version, 'data': data, 'schema': schema,
                    'path': str(self.path), 'modified': datetime.fromtimestamp(self.path.stat().st_mtime).isoformat(),
                    'backupCount': len(list(self.backups.glob('*.xlsx'))) if self.backups.exists() else 0,
                    'capacity': CAPACITY, 'workCalendar': WORK_CALENDAR, 'businessRows': business_rows(data)}

    def validate(self, sheet, record, state):
        required = {'公司': ['公司名称'], '联系人': ['公司编号', '姓名'], '用胶项目': ['公司编号'],
                    '跟进待办': ['公司编号', '下一步具体动作'], '沟通记录': ['公司编号', '日期'],
                    '样品测试': ['项目编号'], '技术经验': ['客户问题原话'], '每日复盘': ['日期'], '测试轮次': ['项目编号', '样品编号', '日期', '状态']}
        for f in required[sheet]:
            if record.get(f, '') == '':
                raise UserError(f'请填写{f}。')
        maps = {n: {r['_id']: r for r in rows} for n, rows in state['data'].items()}
        for name, rows in state['data'].items():
            if len(maps[name]) != len(rows):
                raise UserError(f'{name}存在重复编号，请先在Excel中更正。')
        if sheet == '公司' and any(r['公司名称'].strip().casefold() == record['公司名称'].strip().casefold()
                                  and r['_id'] != record['_id'] for r in state['data']['公司']):
            raise UserError('已存在同名公司，请编辑已有公司。')
        for f, target in RELATIONS.get(sheet, {}).items():
            if record.get(f) and record[f] not in maps[target]:
                raise UserError(f'{f}关联的记录不存在，请重新选择。')
        project = maps['用胶项目'].get(record.get('项目编号'))
        company = record.get('公司编号') or (project or {}).get('公司编号')
        if project and company != project['公司编号']:
            raise UserError('所选项目不属于当前公司。')
        for f in ('联系人编号', '主要联系人编号', '收件联系人编号', '测试负责人编号'):
            person = maps['联系人'].get(record.get(f))
            if person and company and person['公司编号'] != company:
                raise UserError(f'{f}不属于当前公司，请重新选择。')
        for f, target in RELATIONS.get(sheet, {}).items():
            linked = maps[target].get(record.get(f))
            if not linked or target in ('公司', '联系人', '用胶项目'):
                continue
            linked_project = maps['用胶项目'].get(linked.get('项目编号'))
            linked_company = linked.get('公司编号') or (linked_project or {}).get('公司编号')
            if company and linked_company and company != linked_company:
                raise UserError(f'{f}属于另一家公司。')
            if record.get('项目编号') and linked.get('项目编号') and record['项目编号'] != linked['项目编号']:
                raise UserError(f'{f}属于另一个用胶项目。')
        if sheet == '跟进待办' and record.get('状态') == '已完成':
            if not record.get('完成日期') or not record.get('完成证据或反馈'):
                raise UserError('完成任务需要填写完成日期和完成证据。')
        if sheet == '样品测试' and record.get('状态') in SHIPPED:
            if not all(record.get(x) for x in ('完整型号', '数量', '单位', '寄出日期')):
                raise UserError('已寄出的样品需要填写完整型号、数量、单位和寄出日期。')
        if sheet == '测试轮次':
            if record.get('状态') in ('测试中', '通过', '不通过', '待复测'):
                if not all(record.get(f) for f in ('测试负责人编号', '测试条件', '合格标准')):
                    raise UserError('开始测试需要填写测试负责人、测试条件和合格标准。')
            if record.get('状态') in ('通过', '不通过', '待复测'):
                if not record.get('测试结果与原因') or not record.get('证据路径或链接'):
                    raise UserError('记录测试结论需要填写结果与原因，以及证据路径或链接。')
        if sheet in ('样品测试', '测试轮次') and record.get('状态') in ('已寄出', '已签收', '测试中', '待复测'):
            if not record.get('约定反馈日期') or record.get('日期依据') not in ('客户约定', '本人计划'):
                raise UserError('请填写回访或测试反馈日期，并确认日期依据（客户约定或本人计划）。')
        if sheet == '技术经验' and record.get('证据状态') == '同事经验':
            if not record.get('为什么选它') or not record.get('答复人'):
                raise UserError('标成同事经验时，要写下答复人和为什么选这支胶。')

    def save(self, body):
        if body.get('sheet') in ('样品测试', '测试轮次', '跟进待办'):
            result = self.save_batch({'version': body.get('version'), 'records': [body]})
            return dict(result, id=result['ids'][0])
        return self._save(body)

    def transcript_preview(self, body):
        if not isinstance(body.get('id'), str):
            raise UserError('请选择要更新的拜访记录。')
        _, state = self.read()
        record = next((r for r in state['data']['沟通记录'] if r['_id'] == body['id']), None)
        if not record:
            raise UserError('拜访记录已不存在，请刷新。', 409)
        raw, turns = transcript_upload(body, record)
        return {'id': record['_id'], 'date': record['日期'], 'count': len(turns),
                'filename': Path(body['filename']).name, 'sha256': hashlib.sha256(raw).hexdigest(),
                'sample': [{'time': t['time'], 'speaker': t['speaker'], 'text': t['text'][:160]} for t in turns[:3]]}

    def save_transcript(self, body):
        if not isinstance(body.get('id'), str):
            raise UserError('请选择要更新的拜访记录。')
        _, state = self.read()
        record = next((r for r in state['data']['沟通记录'] if r['_id'] == body['id']), None)
        if not record:
            raise UserError('拜访记录已不存在，请刷新。', 409)
        if body.get('version') != state['version']:
            raise UserError('Excel 已变化，请刷新页面后重新预览。', 409)
        raw, turns = transcript_upload(body, record)
        stem = f"{record['_id']}_校对稿_{hashlib.sha256(raw).hexdigest()[:12]}"
        source_dir = originals_dir() / '校对稿'
        clean_dir = originals_dir() / '整理稿'
        source_dir.mkdir(parents=True, exist_ok=True)
        clean_dir.mkdir(parents=True, exist_ok=True)
        source = source_dir / f'{stem}.docx'
        clean = clean_dir / f'{stem}.md'
        lines = [f"# {record['日期']} 拜访校对稿 · {record['_id']}", '',
                 f"- 用户上传文件：{Path(body['filename']).name}",
                 f"- 本地校对稿：`{source}`", f"- SHA-256：`{hashlib.sha256(raw).hexdigest()}`",
                 '- 说明：由用户核对稿直接导入；原始转写字段及原件未覆盖。', '', '## 对话', '']
        for turn in turns:
            lines.extend([f"**{turn['time']} · {turn['speaker']} · 原稿段落 {turn['paragraph']}**", '', turn['text'], ''])
        side_text = lambda side: '\n'.join(f"{t['time']} {t['speaker']} [段落{t['paragraph']}]：{t['text']}" for t in turns if t['side'] == side)
        created = []
        try:
            if not source.exists():
                source.write_bytes(raw)
                created.append(source)
            if not clean.exists():
                clean.write_text('\n'.join(lines), encoding='utf-8')
                created.append(clean)
            result = self.save({'sheet': '沟通记录', 'id': record['_id'], 'version': body['version'],
                                'values': {'客户整理稿及时间戳': side_text('customer'),
                                           '我方整理稿及时间戳': side_text('ours'),
                                           '整理稿路径': str(clean)}})
        except Exception:
            for path in created:
                path.unlink(missing_ok=True)
            raise
        return dict(result, count=len(turns), source=str(source), cleaned=str(clean))

    def _save(self, body):
        with self.lock:
            if (self.path.parent / ('~$' + self.path.name)).exists() or (self.path.parent / ('.~lock.' + self.path.name + '#')).exists():
                raise UserError('请先关闭正在编辑此文件的Excel或LibreOffice，再保存。', 409)
            wb, state = self.read()
            if body.get('version') != state['version']:
                raise UserError('Excel已被其他窗口或程序修改。请保留输入，关闭表单后刷新，再重新编辑。', 409)
            sheet = body.get('sheet')
            if sheet not in SHEETS:
                raise UserError('未知数据表。')
            rows = state['data'][sheet]
            old = next((r for r in rows if r['_id'] == body.get('id')), None)
            if body.get('id') and not old:
                raise UserError('记录已不存在，请刷新。', 409)
            values = body.get('values', {})
            if not isinstance(values, dict):
                raise UserError('记录格式不正确。')
            fields = state['schema'][sheet]
            editable = {f['name']: f for f in fields if not f['computed']}
            id_field = fields[0]['name']
            if set(values) - set(editable) or id_field in values or '日期调整记录' in values or ('推进记录' in values and not body.get('_linked')) or (sheet in ('样品测试', '测试轮次') and '任务编号' in values and not body.get('_linked')):
                raise UserError('编号和自动计算字段不能修改。')
            used = {r['_id'] for r in rows}
            seq = 1
            while f'{SHEETS[sheet]}{seq:04d}' in used:
                seq += 1
            rid = old['_id'] if old else f'{SHEETS[sheet]}{seq:04d}'
            record = {f: (old or {}).get(f, '') for f in editable}
            for f, value in values.items():
                if not isinstance(value, (str, int, float)) or isinstance(value, bool):
                    raise UserError(f'{f}的格式不正确。')
                if isinstance(value, str):
                    value = value.strip()
                    if len(value) > 30000:
                        raise UserError(f'{f}过长，请将完整记录保存在原文件。')
                meta = editable[f]
                if meta['type'] == 'text':
                    value = str(value)
                if value != '' and meta['type'] == 'date':
                    try:
                        value = date.fromisoformat(str(value)).isoformat()
                    except (ValueError, TypeError):
                        raise UserError(f'{f}需要有效日期。')
                if value != '' and meta['type'] == 'number':
                    try:
                        value = float(value)
                    except (ValueError, TypeError):
                        raise UserError(f'{f}需要数字。')
                    if not math.isfinite(value) or value < 0 or (f == '数量' and value == 0):
                        raise UserError(f'{f}需要大于零的数字。' if f == '数量' else f'{f}需要不小于零的数字。')
                if meta['options'] and value and value not in meta['options']:
                    raise UserError(f'{f}请选择列表中的值。')
                record[f] = value
            if sheet == '样品测试' and record.get('签收日期') and not record.get('计划测试日期'):
                record['计划测试日期'] = third_workday(record['签收日期'])
            if sheet == '样品测试' and record.get('计划测试日期') and not record.get('约定反馈日期'):
                record['约定反馈日期'] = record['计划测试日期']
                record['日期依据'] = '本人计划'
            if sheet == '用胶项目':
                if '跟进历史' in values:
                    raise UserError('跟进历史由保存时自动记录。')
                changed_feedback = record.get('最新反馈') and any(record.get(k) != (old or {}).get(k, '') for k in ('最新反馈', '反馈日期'))
                if changed_feedback:
                    if not record.get('反馈日期'):
                        record['反馈日期'] = date.today().isoformat()
                    entry = f"{record['反馈日期']}：{record['最新反馈']}"
                    record['跟进历史'] = '\n'.join(filter(None, [(old or {}).get('跟进历史'), entry]))
            if sheet == '跟进待办' and not body.get('_linked') and '下一步具体动作' in values:
                if values['下一步具体动作'] != (old or {}).get('下一步具体动作'):
                    record['动作来源'] = '手动'
            record[id_field] = rid
            record['_id'] = rid
            if old and any(old.get(k, '') != record.get(k, '') for k in ('公司编号', '项目编号')):
                for source, links in RELATIONS.items():
                    for f, target in links.items():
                        if target == sheet and any(r.get(f) == rid for r in state['data'][source]):
                            raise UserError('此记录已有下游关联，不能直接更换所属公司或项目。请新建对应记录并逐项调整关联。')
            if sheet == '跟进待办' and old and old.get('计划日期') and old.get('计划日期') != record.get('计划日期'):
                if not record.get('延期原因'):
                    raise UserError('调整已有计划日期需要填写延期原因，原日期会保留。')
                entry = f"{date.today().isoformat()}：{old['计划日期']} → {record.get('计划日期') or '待定'}；{record['延期原因']}"
                record['日期调整记录'] = '\n'.join(filter(None, [old.get('日期调整记录'), entry]))
            self.validate(sheet, record, state)
            s = wb[sheet]
            row = old['_row'] if old else next((r for r in range(8, LAST_ROW + 1) if s.cell(r, 1).value is None), None)
            if row is None:
                raise UserError(f'本表已达{CAPACITY}条，请先扩展工作簿容量。')
            for i, field in enumerate(fields, 1):
                if field['computed']:
                    if not old and s.cell(8, i).data_type == 'f':
                        s.cell(row, i).value = Translator(s.cell(8, i).value, origin=s.cell(8, i).coordinate).translate_formula(s.cell(row, i).coordinate)
                        s.cell(row, i)._style = copy(s.cell(8, i)._style)
                    continue
                v = record[field['name']]
                c = s.cell(row, i)
                if not old:
                    c._style = copy(s.cell(8, i)._style)
                if v == '':
                    c.value = None
                elif field['type'] == 'date':
                    c.value = date.fromisoformat(v)
                    c.number_format = 'yyyy-mm-dd'
                else:
                    c.value = v
                    if isinstance(v, str):
                        c.data_type = 's'
            wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
            fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
            os.close(fd)
            try:
                refresh_business_sheet(wb)
                wb.save(temp)
                test = load_workbook(temp)
                if test[sheet].cell(row, 1).value != rid:
                    raise UserError('保存校验失败，原文件保持不变。', 500)
                if digest(self.path) != state['version']:
                    raise UserError('保存期间Excel发生变化，请刷新后重试。', 409)
                self.backups.mkdir(exist_ok=True)
                backup = self.backups / f'{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx'
                shutil.copy2(self.path, backup)
                os.replace(temp, self.path)
            finally:
                if os.path.exists(temp):
                    os.unlink(temp)
            return {'id': rid, 'version': digest(self.path), 'backup': str(backup)}


    def save_batch(self, body):
        records = body.get('records')
        if not isinstance(records, list) or not 1 <= len(records) <= 50 or not all(isinstance(r, dict) for r in records):
            raise UserError('请提交 1 至 50 条有效记录。')
        with self.lock:
            version = digest(self.path)
            if body.get('version') != version:
                raise UserError('Excel已被其他窗口或程序修改，请刷新。', 409)
            if any((self.path.parent / n).exists() for n in ('~$' + self.path.name, '.~lock.' + self.path.name + '#')):
                raise UserError('请先关闭正在编辑此文件的Excel或LibreOffice，再保存。', 409)
            with tempfile.TemporaryDirectory(dir=self.path.parent) as folder:
                target = Path(folder) / 'batch.xlsx'
                shutil.copy2(self.path, target)
                staged = Store(target)
                ids = {}
                results = []
                for item in records:
                    values = item.get('values')
                    if not isinstance(values, dict):
                        raise UserError('记录格式不正确。')
                    values = dict(values)
                    for field, value in values.items():
                        if isinstance(value, str) and value.startswith('@record:'):
                            key = value[8:]
                            if key not in ids:
                                raise UserError('关联记录顺序不正确。')
                            values[field] = ids[key]
                    clean = {k: v for k, v in item.items() if not k.startswith('_')}
                    previous = next((r for r in staged.read()[1]['data'].get(item.get('sheet'), []) if r['_id'] == item.get('id')), None)
                    result = staged._save(dict(clean, values=values, version=digest(target)))
                    staged.sync_followup(item.get('sheet'), result['id'], previous)
                    if item.get('key'):
                        ids[item['key']] = result['id']
                    results.append(result['id'])
                if digest(self.path) != version:
                    raise UserError('保存期间Excel发生变化，请刷新后重试。', 409)
                self.backups.mkdir(exist_ok=True)
                backup = self.backups / f'{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx'
                shutil.copy2(self.path, backup)
                os.replace(target, self.path)
                return {'ids': results, 'version': digest(self.path), 'backup': str(backup)}


    def sync_followup(self, sheet, rid, previous):
        state = self.read()[1]
        current = next(r for r in state['data'][sheet] if r['_id'] == rid)
        def write(name, values, ident=''):
            return self._save({'sheet': name, 'id': ident, 'values': values, 'version': digest(self.path), '_linked': True})
        if sheet == '跟进待办':
            # The dated task is the same appointment shown on its sample/trial.
            for name in ('样品测试', '测试轮次'):
                for source in state['data'][name]:
                    if source.get('任务编号') == rid and previous and current.get('计划日期') != previous.get('计划日期'):
                        entry = f"{date.today().isoformat()}：跟进日期 {previous.get('计划日期') or '待定'} → {current.get('计划日期') or '待定'}；{current.get('延期原因') or '补充计划'}"
                        write(name, {'约定反馈日期': current['计划日期'], '日期依据': current.get('日期依据') or '待确认', '推进记录': '\n'.join(filter(None, [source.get('推进记录'), entry]))}, source['_id'])
            return
        if sheet not in ('样品测试', '测试轮次'):
            return
        old = previous or {}
        changed = not previous or any(current.get(k) != old.get(k) for k in ('状态', '约定反馈日期', '计划测试日期', '测试结果与原因', '问题与下一步', '下一步具体动作'))
        if changed:
            entry = f"{date.today().isoformat()}：登记{sheet} {current.get('状态') or '待确认'}；回访/反馈 {current.get('约定反馈日期') or '待定'}"
            if current.get('计划测试日期'):
                entry += '；计划测试 ' + current['计划测试日期']
            reason = current.get('测试结果与原因') or current.get('问题与下一步') or current.get('下一步具体动作')
            if reason:
                entry += '；' + reason
            write(sheet, {'推进记录': '\n'.join(filter(None, [current.get('推进记录'), entry]))}, rid)
        task = next((r for r in state['data']['跟进待办'] if r['_id'] == current.get('任务编号')), None)
        if not changed and task:
            return
        if task and task['状态'] in ('已完成', '已取消') and not changed:
            return
        step = STEP_ACTIONS.get(current.get('状态') or '')
        kind, verb = step or ('测试反馈', '确认下一步')
        model = current.get('完整型号') or current.get('物料名称') or rid
        auto_action = f'{model} · {verb}'
        manual = bool(task and task.get('动作来源') == '手动' and task.get('下一步具体动作'))
        project = next(r for r in state['data']['用胶项目'] if r['_id'] == current['项目编号'])
        sample_id = rid if sheet == '样品测试' else current['样品编号']
        values = {'公司编号': project['公司编号'], '项目编号': project['_id'], '样品编号': sample_id,
                  '事项类型': task.get('事项类型') if manual else kind,
                  '下一步具体动作': task.get('下一步具体动作') if manual else auto_action,
                  '动作来源': '手动' if manual else ('样品自动' if sheet == '样品测试' else '测试自动'),
                  '自动步骤': '' if manual else (current.get('状态') or ''),
                  '负责人': (task or {}).get('负责人') or '我',
                  '计划日期': current.get('约定反馈日期') or '',
                  '日期依据': current.get('日期依据') if current.get('日期依据') in ('客户约定', '本人计划') else ''}
        if task and task['状态'] in ('已完成', '已取消'):
            task = None
        if not task:
            values['状态'] = '待办'
        elif task.get('计划日期') and values['计划日期'] != task['计划日期']:
            if not current.get('延期原因'):
                raise UserError('调整样品或测试的回访日期，请填写延期原因，关联待办将同步调整。')
            values['延期原因'] = current['延期原因']
        result = write('跟进待办', values, task['_id'] if task else '')
        write(sheet, {'任务编号': result['id']}, rid)


    def save_project(self, body):
        pid = body.get('id')
        if not isinstance(pid, str) or not pid:
            raise UserError('请选择项目。')
        project_values = body.get('project') or {}
        samples = body.get('samples') or []
        follow = body.get('follow') or {}
        if not isinstance(project_values, dict) or not isinstance(samples, list) or not isinstance(follow, dict):
            raise UserError('项目格式不正确。')
        allowed = {'最新反馈', '反馈日期', '需要得到的支持', '客户问题与目标', '用胶部位', '项目备注', '当前工艺与固化条件', '现用材料或型号', '客户生产用量', '用量单位及周期', '开始日期', '关闭日期', '应用端客户名称', '项目发起人', '业务承接人', '项目开放原因', '竞品主要问题', '包装规格', '当前售价', '可接受价格预估', '项目外部现状', '方案性质'}
        sample_fields = {'状态', '签收日期', '计划测试日期', '寄出日期', '约定反馈日期', '日期依据', '延期原因', '测试结果与证据', '快递单号', '数量', '单位', '问题与下一步', '备注'}
        with self.lock:
            version = digest(self.path)
            if body.get('version') != version:
                raise UserError('Excel已被其他窗口或程序修改，请刷新。', 409)
            if any((self.path.parent / n).exists() for n in ('~$' + self.path.name, '.~lock.' + self.path.name + '#')):
                raise UserError('请先关闭正在编辑此文件的Excel或LibreOffice，再保存。', 409)
            with tempfile.TemporaryDirectory(dir=self.path.parent) as folder:
                target = Path(folder) / 'project.xlsx'
                shutil.copy2(self.path, target)
                staged = Store(target)
                state = staged.read()[1]
                project = next((r for r in state['data']['用胶项目'] if r['_id'] == pid), None)
                if not project:
                    raise UserError('项目已不存在，请刷新。', 409)
                values = {k: v for k, v in project_values.items() if k in allowed}
                if values:
                    staged._save({'sheet': '用胶项目', 'id': pid, 'values': values, 'version': digest(target)})
                for item in samples:
                    if not isinstance(item, dict) or not isinstance(item.get('id'), str):
                        raise UserError('请选择要更新的样品。')
                    current = next((r for r in staged.read()[1]['data']['样品测试'] if r['_id'] == item['id']), None)
                    if not current or current.get('项目编号') != pid:
                        raise UserError('样品不属于这个项目。')
                    migration = str(item.get('迁移记录') or '').strip()
                    if migration:
                        history = '\n'.join(filter(None, [current.get('推进记录'), '迁移记录：' + migration]))
                        staged._save({'sheet': '样品测试', 'id': item['id'], 'values': {'推进记录': history}, 'version': digest(target), '_linked': True})
                        current = next(r for r in staged.read()[1]['data']['样品测试'] if r['_id'] == item['id'])
                    sample_values = {k: v for k, v in item.items() if k in sample_fields}
                    if not sample_values:
                        continue
                    saved = staged._save({'sheet': '样品测试', 'id': item['id'], 'values': sample_values, 'version': digest(target)})
                    staged.sync_followup('样品测试', saved['id'], current)
                action = str(follow.get('下一步具体动作') or '').strip()
                if action:
                    tasks = [t for t in staged.read()[1]['data']['跟进待办'] if t.get('项目编号') == pid and t.get('状态') not in ('已完成', '已取消')]
                    tasks.sort(key=lambda r: (r.get('计划日期') or '9999-12-31', r['_id']))
                    chosen = next((t for t in tasks if t['_id'] == follow.get('taskId')), None) if follow.get('taskId') else (tasks[0] if tasks else None)
                    plan = str(follow.get('计划日期') or '')
                    if not chosen or (chosen.get('下一步具体动作') or '') != action or (chosen.get('计划日期') or '') != plan:
                        follow_values = {'下一步具体动作': action, '计划日期': plan, '日期依据': follow.get('日期依据') or ('本人计划' if plan else '')}
                        if chosen and chosen.get('计划日期') and plan != (chosen.get('计划日期') or ''):
                            if not str(follow.get('延期原因') or '').strip():
                                raise UserError('调整已有计划日期需要填写延期原因，原日期会保留。')
                            follow_values['延期原因'] = str(follow['延期原因'])
                        if not chosen:
                            follow_values.update({'公司编号': project['公司编号'], '项目编号': pid, '事项类型': '定期回访', '状态': '待办', '负责人': '我'})
                        staged._save({'sheet': '跟进待办', 'id': chosen['_id'] if chosen else '', 'values': follow_values, 'version': digest(target)})
                if digest(self.path) != version:
                    raise UserError('保存期间Excel发生变化，请刷新后重试。', 409)
                self.backups.mkdir(exist_ok=True)
                backup = self.backups / f'{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx'
                shutil.copy2(self.path, backup)
                os.replace(target, self.path)
                return {'id': pid, 'version': digest(self.path), 'backup': str(backup)}


    def delete_sample(self, body):
        sheet_name = body.get('sheet')
        if sheet_name not in ('样品测试', '跟进待办', '用胶项目'):
            raise UserError('目前只能删除项目、样品或跟进。')
        rid = body.get('id')
        if not isinstance(rid, str) or not rid:
            raise UserError('请选择要删除的记录。')
        with self.lock:
            if body.get('version') != digest(self.path):
                raise UserError('Excel已被其他窗口或程序修改。请刷新后再删除。', 409)
            if any((self.path.parent / n).exists() for n in ('~$' + self.path.name, '.~lock.' + self.path.name + '#')):
                raise UserError('请先关闭正在编辑此文件的Excel或LibreOffice，再删除。', 409)
            wb, state = self.read()
            target = next((r for r in state['data'][sheet_name] if r['_id'] == rid), None)
            if not target:
                raise UserError('记录已不存在，请刷新。', 409)
            if sheet_name == '用胶项目':
                samples = {r['_id'] for r in state['data']['样品测试'] if r.get('项目编号') == rid}
                removed = {
                    '用胶项目': {rid},
                    '样品测试': samples,
                    '测试轮次': {r['_id'] for r in state['data']['测试轮次'] if r.get('项目编号') == rid or r.get('样品编号') in samples},
                    '跟进待办': {r['_id'] for r in state['data']['跟进待办'] if r.get('项目编号') == rid or r.get('样品编号') in samples},
                }
                for source, links in RELATIONS.items():
                    ws = wb[source]
                    headers = [c.value for c in ws[7]]
                    for item in state['data'][source]:
                        if item['_id'] in removed.get(source, set()):
                            continue
                        for field, linked_sheet in links.items():
                            if item.get(field) in removed.get(linked_sheet, set()):
                                ws.cell(item['_row'], headers.index(field) + 1).value = None
                for name, ids in removed.items():
                    for item in state['data'][name]:
                        if item['_id'] not in ids:
                            continue
                        for col, field in enumerate(state['schema'][name], 1):
                            cell = wb[name].cell(item['_row'], col)
                            if not field['computed'] and cell.data_type != 'f':
                                cell.value = None
                row = target['_row']
            else:
                if sheet_name == '样品测试' and any(r.get('样品编号') == rid for r in state['data'].get('测试轮次', [])):
                    raise UserError('这款样品已有测试轮次，请先处理轮次后再删除。')
                links = [('跟进待办', '样品编号')] if sheet_name == '样品测试' else [('样品测试', '任务编号'), ('测试轮次', '任务编号'), ('沟通记录', '下一步任务编号')]
                for link_sheet, link_field in links:
                    if link_sheet not in state['data']:
                        continue
                    ws = wb[link_sheet]
                    col = [c.value for c in ws[7]].index(link_field) + 1
                    for row_item in state['data'][link_sheet]:
                        if row_item.get(link_field) == rid:
                            ws.cell(row_item['_row'], col).value = None
                sheet = wb[sheet_name]
                row = target['_row']
                for i, field in enumerate(state['schema'][sheet_name], 1):
                    cell = sheet.cell(row, i)
                    if field['computed'] or cell.data_type == 'f':
                        continue
                    cell.value = None
            wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
            fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
            os.close(fd)
            try:
                refresh_business_sheet(wb)
                wb.save(temp)
                test = load_workbook(temp)
                if test[sheet_name].cell(row, 1).value is not None:
                    raise UserError('删除校验失败，原文件保持不变。', 500)
                if digest(self.path) != body.get('version'):
                    raise UserError('删除期间Excel发生变化，请刷新后重试。', 409)
                self.backups.mkdir(exist_ok=True)
                backup = self.backups / f'{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx'
                shutil.copy2(self.path, backup)
                os.replace(temp, self.path)
            finally:
                if os.path.exists(temp):
                    os.unlink(temp)
            return {'id': rid, 'version': digest(self.path), 'backup': str(backup)}


def handler_for(store, token):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass

        def send(self, status, content, content_type='application/json; charset=utf-8'):
            if isinstance(content, dict):
                content = json.dumps(content, ensure_ascii=False).encode()
            if isinstance(content, str):
                content = content.encode()
            self.send_response(status)
            self.send_header('Content-Type', content_type)
            self.send_header('Content-Length', str(len(content)))
            self.send_header('Cache-Control', 'no-store')
            self.send_header('X-Content-Type-Options', 'nosniff')
            self.send_header('X-Frame-Options', 'DENY')
            self.end_headers()
            self.wfile.write(content)

        def trusted(self):
            host = self.headers.get('Host', '')
            if host not in (f'127.0.0.1:{self.server.server_port}', f'localhost:{self.server.server_port}'):
                raise UserError('仅支持本机访问。', 403)
            origin = self.headers.get('Origin')
            if origin and origin != 'http://' + host:
                raise UserError('不接受其他网站的请求。', 403)

        def do_GET(self):
            try:
                self.trusted()
                route = urlparse(self.path).path
                if route == '/api/state':
                    with store.lock:
                        _, state = store.read()
                    self.send(200, state)
                elif route == '/api/strategy':
                    folder = ROOT / 'data' / '战略日报'
                    reports = [{'date': f.stem, 'name': f.name} for f in sorted(folder.glob('????-??-??.md'), reverse=True) if re.fullmatch(r'\d{4}-\d{2}-\d{2}', f.stem)]
                    self.send(200, {'reports': reports})
                elif route.startswith('/strategy/'):
                    name = unquote(route.removeprefix('/strategy/'))
                    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}\.(md|png|html)', name):
                        self.send(404, {'error': '日报不存在。'})
                        return
                    file = ROOT / 'data' / '战略日报' / name
                    if not file.is_file():
                        self.send(404, {'error': '日报不存在。'})
                        return
                    content_type = {'.md': 'text/plain; charset=utf-8', '.png': 'image/png', '.html': 'text/html; charset=utf-8'}[file.suffix]
                    self.send(200, file.read_bytes(), content_type)
                elif route == '/download':
                    self.send(200, store.path.read_bytes(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                elif route in ('/', '/app.js', '/style.css'):
                    f = {'/': 'index.html', '/app.js': 'app.js', '/style.css': 'style.css'}[route]
                    content = (ROOT / 'dist' / f).read_text().replace('__TOKEN__', token)
                    typ = {'/': 'text/html', '/app.js': 'text/javascript', '/style.css': 'text/css'}[route]
                    self.send(200, content, typ + '; charset=utf-8')
                else:
                    self.send(404, {'error': '页面不存在。'})
            except UserError as e:
                self.send(e.status, {'error': str(e)})
            except (OSError, KeyError, ValueError):
                self.send(500, {'error': '无法读取工作簿，请检查文件是否完整，或关闭Excel后重试。'})

        def do_POST(self):
            try:
                self.trusted()
                if self.headers.get('X-Workspace-Token') != token:
                    raise UserError('页面已失效，请刷新。', 403)
                if self.path not in ('/api/records', '/api/batch', '/api/delete', '/api/transcript/preview', '/api/transcript/save', '/api/project-update'):
                    raise UserError('接口不存在。', 404)
                length = int(self.headers.get('Content-Length', 0))
                limit = 2200000 if self.path.startswith('/api/transcript/') else 1000000
                if not 0 < length <= limit:
                    raise UserError('提交内容为空或过大。')
                body = json.loads(self.rfile.read(length))
                if not isinstance(body, dict):
                    raise UserError('请求格式错误。')
                result = (store.save_project(body) if self.path == '/api/project-update' else
                          store.transcript_preview(body) if self.path == '/api/transcript/preview' else
                          store.save_transcript(body) if self.path == '/api/transcript/save' else
                          store.delete_sample(body) if self.path == '/api/delete' else
                          store.save_batch(body) if self.path == '/api/batch' else store.save(body))
                self.send(200, result)
            except UserError as e:
                self.send(e.status, {'error': str(e)})
            except (ValueError, TypeError):
                self.send(400, {'error': '请求格式错误。'})
            except OSError:
                self.send(500, {'error': '文件保存失败，请检查磁盘空间及文件权限。'})
    return Handler


def personal_root():
    here = ROOT.parent
    if (here / '个人客户数据').is_dir() or (here / '客户数据原稿').is_dir():
        return here
    return None


def default_data_path():
    root = personal_root()
    if root:
        return root / '个人客户数据' / '客户数据.xlsx'
    return ROOT / 'data' / '客户数据.xlsx'


def originals_dir():
    root = personal_root()
    if root:
        return root / '客户数据原稿'
    return ROOT / '原稿'


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--port', type=int, default=8765)
    p.add_argument('--data', type=Path, default=None)
    p.add_argument('--open', action='store_true')
    args = p.parse_args()
    store = Store(args.data or default_data_path())
    try:
        server = ThreadingHTTPServer(('127.0.0.1', args.port), handler_for(store, secrets.token_urlsafe(32)))
    except OSError:
        print(f'端口{args.port}已被占用。若工作台已启动，请在浏览器打开 http://127.0.0.1:{args.port}；否则先关闭占用端口的程序。', flush=True)
        return
    url = f'http://127.0.0.1:{server.server_port}'
    print(f'销售工作台：{url}\nExcel：{store.path}\n关闭此窗口将停止工作台。', flush=True)
    if args.open:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == '__main__':
    main()
