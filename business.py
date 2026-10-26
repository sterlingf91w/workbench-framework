"""One projection shared by the business page and the Excel front sheet."""
from collections import Counter
from copy import copy
from datetime import date, datetime
import math

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ['编号', '开始日期', '关闭日期', '应用端客户名称', '项目发起人', '项目联系人（职位）',
           '客户联系方式', '业务承接人', '用胶点及要求', '客户用胶工艺', '项目开放原因',
           '竞品信息', '竞品主要问题', '包装规格', '月用量（单位）', '当前售价', '可接受价格预估',
           '项目外部现状', '推荐方案', '方案性质', '项目内部状态', '需要得到的支持', '下步计划/反馈', '备注']
STATUS_ORDER = ['待确认', '已申请', '待寄出', '已寄出', '已签收', '测试中', '通过', '不通过', '暂停', '取消']
STEP_ACTIONS = {'已申请': ('送样', '跟进寄出'), '待寄出': ('送样', '安排寄出'),
                '已寄出': ('确认签收', '确认收样'), '已签收': ('测试反馈', '安排测试'),
                '测试中': ('测试反馈', '收集反馈'), '待复测': ('测试反馈', '安排复测'),
                '通过': ('测试反馈', '确认测试结论'), '不通过': ('测试反馈', '确认原因与下一步')}


def sample_label(status):
    return '已收样' if status == '已签收' else status or '待确认'


def business_rows(data):
    companies = {r['_id']: r for r in data.get('公司', [])}
    people = {r['_id']: r for r in data.get('联系人', [])}
    result = []
    for project in data.get('用胶项目', []):
        pid = project['_id']
        company = companies.get(project.get('公司编号'), {})
        samples = [r for r in data.get('样品测试', []) if r.get('项目编号') == pid]
        tasks = [r for r in data.get('跟进待办', []) if r.get('项目编号') == pid and r.get('状态') not in ('已完成', '已取消')]
        tasks.sort(key=lambda r: (r.get('计划日期') or '9999-12-31', r['_id']))
        primary = tasks[0] if tasks else {}
        person = people.get(project.get('主要联系人编号'), {})
        if not person:
            candidates = {s.get('收件联系人编号') for s in samples if s.get('收件联系人编号')}
            if len(candidates) == 1:
                person = people.get(next(iter(candidates)), {})
        counts = Counter(s.get('状态') or '待确认' for s in samples)
        summary = '、'.join(f'{counts[k]}款{sample_label(k)}' for k in STATUS_ORDER if counts[k]) or '尚无样品'
        models = list(dict.fromkeys(s.get('完整型号') or s.get('物料名称') for s in samples if s.get('完整型号') or s.get('物料名称')))
        if project.get('候选型号') and project['候选型号'] not in models:
            models.append(project['候选型号'])
        feedback = project.get('最新反馈', '')
        feedback_date = project.get('反馈日期', '')
        next_lines = [' · '.join(filter(None, [t.get('计划日期') or '日期待定', t.get('下一步具体动作')])) for t in tasks]
        detail = '\n'.join(f"{s.get('完整型号') or s.get('物料名称') or s['_id']}：{sample_label(s.get('状态'))}" for s in samples)
        internal = '\n'.join(filter(None, [detail or summary, (feedback_date + ' ' + feedback).strip() if feedback else '', project.get('当前卡点', '')]))
        contact = '、'.join(filter(None, [person.get('电话'), person.get('邮箱')]))
        person_text = person.get('姓名', '') + (f"（{person['职位']}）" if person.get('职位') else '')
        glue = '；'.join(filter(None, [project.get('用胶部位'), project.get('客户问题与目标')]))
        usage = ' '.join(str(x) for x in [project.get('客户生产用量'), project.get('用量单位及周期')] if x != '' and x is not None)
        packing = project.get('包装规格') or '；'.join(dict.fromkeys(s['规格'] for s in samples if s.get('规格')))
        cells = [pid, project.get('开始日期', ''), project.get('关闭日期', ''), project.get('应用端客户名称') or company.get('公司名称', ''),
                 project.get('项目发起人', ''), person_text, contact, project.get('业务承接人') or company.get('客户负责人', ''),
                 glue, project.get('当前工艺与固化条件', ''), project.get('项目开放原因', ''), project.get('现用材料或型号', ''),
                 project.get('竞品主要问题', ''), packing, usage, project.get('当前售价', ''), project.get('可接受价格预估', ''),
                 project.get('项目外部现状', ''), '\n'.join(models), project.get('方案性质', ''), internal,
                 project.get('需要得到的支持', ''), '\n'.join(next_lines), project.get('项目备注', '')]
        result.append({'id': pid, 'companyId': company.get('_id', ''), 'company': company.get('公司名称', ''),
                       'application': project.get('终端产品', ''), 'requirements': glue, 'person': person_text,
                       'models': models, 'sampleSummary': summary, 'feedback': feedback, 'feedbackDate': feedback_date,
                       'nextAction': primary.get('下一步具体动作', ''), 'nextDate': primary.get('计划日期', ''),
                       'primaryTaskId': primary.get('_id', ''), 'support': project.get('需要得到的支持', ''),
                       'tasks': tasks, 'samples': samples, 'project': project, 'cells': cells})
    return result


def workbook_data(wb):
    data = {}
    for name in ('公司', '联系人', '用胶项目', '样品测试', '跟进待办', '沟通记录'):
        ws = wb[name]
        headers = [c.value for c in ws[7]]
        data[name] = []
        for row in ws.iter_rows(min_row=8, values_only=True):
            if not row[0]:
                continue
            record = {k: v.strftime('%Y-%m-%d') if isinstance(v, (date, datetime)) else '' if v is None else v
                      for k, v in zip(headers, row) if k and '〔自动〕' not in str(k)}
            record['_id'] = str(row[0])
            data[name].append(record)
    return data


def refresh_business_sheet(wb):
    rows = business_rows(workbook_data(wb))
    ws = wb['业务跟进表'] if '业务跟进表' in wb.sheetnames else wb.create_sheet('业务跟进表', 0)
    wb.move_sheet(ws, offset=-wb.sheetnames.index(ws.title))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws.merge_cells('A1:X1')
    ws['A1'] = '业务项目跟进表 · 网页更新后自动汇总 · 明细与历史保留在后续工作表'
    ws['A1'].font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='183C50')
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 38
    widths = [12, 14, 14, 32, 24, 24, 25, 20, 44, 32, 24, 28, 28, 22, 24, 20, 24, 36, 28, 22, 48, 36, 52, 40]
    for col, heading in enumerate(HEADERS, 1):
        c = ws.cell(2, col, heading)
        c.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='276C79' if col in (19, 21, 22, 23) else '365366')
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.row_dimensions[2].height = 32
    for rownum, item in enumerate(rows, 3):
        line_counts = []
        for col, value in enumerate(item['cells'], 1):
            c = ws.cell(rownum, col, value or None)
            if isinstance(value, str):
                c.data_type = 's'
            if col in (2, 3) and value:
                c.value = date.fromisoformat(value)
                c.number_format = 'yyyy-mm-dd'
            c.font = Font(name='微软雅黑', size=11, color='173747')
            c.fill = PatternFill('solid', fgColor='F0F6F8' if rownum % 2 else 'FFFFFF')
            c.alignment = Alignment(wrap_text=True, vertical='top')
            line_counts.append(sum(max(1, math.ceil(sum(2 if ord(ch) > 127 else 1 for ch in part) / max(8, widths[col-1]-2))) for part in str(value or '').split('\n')))
        ws.row_dimensions[rownum].height = max(72, min(409, max(line_counts) * 17 + 12))
    ws.freeze_panes = 'E3'
    ws.auto_filter.ref = f'A2:X{max(3, len(rows)+2)}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:2'
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f'A1:X{max(3, len(rows)+2)}'
    wb.active = 0
