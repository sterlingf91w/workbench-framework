import importlib.util
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location('anda_server', Path(__file__).with_name('server.py'))
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = app.Store(Path(self.tmp.name) / 'data.xlsx')

    def tearDown(self):
        self.tmp.cleanup()

    def save(self, sheet, values, rid=''):
        _, state = self.store.read()
        return self.store.save({'sheet': sheet, 'values': values, 'id': rid, 'version': state['version']})

    def company(self, name='测试公司'):
        return self.save('公司', {'公司名称': name, '市场': '国内'})['id']

    def test_lists_and_formula_text_stay(self):
        before = load_workbook(self.store.path)
        formulas = {(s.title, c.coordinate): c.value for s in before for row in s for c in row if c.data_type == 'f'}
        channels = before['沟通记录'].data_validations.dataValidation
        self.assertTrue(any('朋友圈' in (dv.formula1 or '') for dv in channels))
        statuses = before['样品测试'].data_validations.dataValidation
        self.assertTrue(any('已申请' in (dv.formula1 or '') for dv in statuses))
        self.company()
        after = load_workbook(self.store.path)
        self.assertEqual(formulas, {(s.title, c.coordinate): c.value for s in after for row in s for c in row if c.data_type == 'f'})

    def test_literal_text_and_duplicate_company(self):
        cid = self.company('=1+1')
        cell = load_workbook(self.store.path)['公司']['B8']
        self.assertEqual(cell.value, '=1+1')
        self.assertEqual(cell.data_type, 's')
        with self.assertRaisesRegex(app.UserError, '同名'):
            self.company('=1+1')
        self.save('公司', {'公司名称': '已更新'}, cid)
        self.assertEqual(self.store.read()[1]['data']['公司'][0]['公司名称'], '已更新')

    def test_stale_version_rejected(self):
        _, old = self.store.read()
        self.company()
        digest = app.digest(self.store.path)
        with self.assertRaisesRegex(app.UserError, '已被其他'):
            self.store.save({'sheet': '公司', 'values': {'公司名称': '另一家公司'}, 'version': old['version']})
        self.assertEqual(digest, app.digest(self.store.path))

    def test_cross_company_relation(self):
        c1 = self.company('甲')
        c2 = self.company('乙')
        person = self.save('联系人', {'公司编号': c1, '姓名': '张工'})['id']
        with self.assertRaisesRegex(app.UserError, '不属于'):
            self.save('用胶项目', {'公司编号': c2, '主要联系人编号': person})

    def test_applied_sample_can_omit_shipment(self):
        cid = self.company()
        project = self.save('用胶项目', {'公司编号': cid, '终端产品': '模块'})['id']
        sid = self.save('样品测试', {'项目编号': project, '状态': '已申请', '完整型号': 'AP-8120'})['id']
        self.assertTrue(sid.startswith('S'))
        with self.assertRaisesRegex(app.UserError, '已寄出'):
            self.save('样品测试', {'状态': '已寄出'}, sid)

    def test_moments_channel_and_colleague_experience(self):
        cid = self.company()
        vid = self.save('沟通记录', {'公司编号': cid, '日期': '2026-10-16', '渠道': '朋友圈', '我的角色': '主谈'})['id']
        self.assertTrue(vid.startswith('V'))
        with self.assertRaisesRegex(app.UserError, '为什么选'):
            self.save('技术经验', {'客户问题原话': '用什么胶', '证据状态': '同事经验', '答复人': '同事'})
        kid = self.save('技术经验', {'客户问题原话': '用什么胶', '证据状态': '同事经验', '答复人': '同事', '为什么选它': '间隙需要填，不承担结构粘接'})['id']
        self.assertTrue(kid.startswith('K'))

    def test_capacity_migration_and_201st_record(self):
        wb = load_workbook(self.store.path)
        ws = wb['公司']
        for row in range(8, 208):
            ws.cell(row, 1, f'C{row-7:04d}')
            ws.cell(row, 2, f'容量验证公司{row-7}')
        wb.save(self.store.path)
        cid = self.company('第201家公司')
        self.assertEqual(cid, 'C0201')
        wb = load_workbook(self.store.path)
        self.assertEqual(wb['公司']['A208'].value, cid)
        self.assertIn('A208', wb['公司']['M208'].value)
        self.assertIn('$A$2007', wb['公司']['M208'].value)
        self.assertEqual(wb['公司'].tables['Anda2'].ref, 'A7:M2007')
        self.assertTrue(any('C2007' in str(d.sqref) for d in wb['公司'].data_validations.dataValidation))
        self.assertEqual(len(app.Store(self.store.path).read()[1]['data']['公司']), 201)

    def test_reschedule_preserves_history(self):
        cid = self.company()
        tid = self.save('跟进待办', {'公司编号': cid, '下一步具体动作': '提供资料', '计划日期': '2026-10-17'})['id']
        with self.assertRaisesRegex(app.UserError, '延期原因'):
            self.save('跟进待办', {'计划日期': '2026-10-20'}, tid)
        self.save('跟进待办', {'计划日期': '2026-10-20', '延期原因': '客户改期'}, tid)
        task = self.store.read()[1]['data']['跟进待办'][0]
        self.assertIn('2026-10-17 → 2026-10-20', task['日期调整记录'])
        with self.assertRaises(app.UserError):
            self.save('跟进待办', {'日期调整记录': '伪造历史'}, tid)

    def test_batch_rollback_links_and_backup_restore(self):
        cid = self.company()
        old_hash = app.digest(self.store.path)
        with self.assertRaises(app.UserError):
            self.store.save_batch({'version': old_hash, 'records': [
                {'sheet': '沟通记录', 'values': {'公司编号': cid, '日期': '2026-10-17'}},
                {'sheet': '跟进待办', 'values': {'公司编号': cid}}
            ]})
        self.assertEqual(old_hash, app.digest(self.store.path))
        result = self.store.save_batch({'version': old_hash, 'records': [
            {'sheet': '沟通记录', 'key': 'talk', 'values': {'公司编号': cid, '日期': '2026-10-17'}},
            {'sheet': '跟进待办', 'values': {'公司编号': cid, '下一步具体动作': '确认标准', '来源沟通编号': '@record:talk'}}
        ]})
        state = app.Store(self.store.path).read()[1]
        self.assertEqual(state['data']['跟进待办'][0]['来源沟通编号'], result['ids'][0])
        self.assertEqual(app.digest(Path(result['backup'])), old_hash)
        import shutil
        shutil.copy2(result['backup'], self.store.path)
        self.assertEqual(app.digest(self.store.path), old_hash)
        self.assertFalse(app.Store(self.store.path).read()[1]['data']['沟通记录'])

    def test_upgrade_preserves_existing_inputs_and_is_idempotent(self):
        import shutil
        source = Path(__file__).with_name('template') / '空白跟进表.xlsx'
        dest = Path(self.tmp.name) / 'real-copy.xlsx'
        shutil.copy2(source, dest)
        before = load_workbook(dest)
        store = app.Store(dest)
        after = load_workbook(dest)
        for name in app.SHEETS:
            if name not in before.sheetnames:
                continue
            for row in before[name].iter_rows(min_row=7):
                for c in row:
                    if c.data_type != 'f':
                        self.assertEqual(c.value, after[name][c.coordinate].value, (name, c.coordinate))
        old_hash = app.digest(dest)
        app.Store(dest)
        self.assertEqual(old_hash, app.digest(dest))
        self.assertEqual(store.read()[1]['capacity'], 2000)

    def test_trial_rounds_keep_history_and_reject_invalid_links(self):
        cid = self.company()
        project = self.save('用胶项目', {'公司编号': cid})['id']
        other = self.save('用胶项目', {'公司编号': cid})['id']
        sample = self.save('样品测试', {'项目编号': project, '完整型号': '试验胶'})['id']
        person = self.save('联系人', {'公司编号': cid, '姓名': '测试工程师'})['id']
        base = {'项目编号': project, '样品编号': sample, '日期': '2026-10-18', '状态': '待安排'}
        with self.assertRaisesRegex(app.UserError, '另一个用胶项目'):
            self.save('测试轮次', dict(base, 项目编号=other))
        with self.assertRaisesRegex(app.UserError, '测试负责人'):
            self.save('测试轮次', dict(base, 状态='测试中'))
        full = dict(base, 测试负责人编号=person, 测试条件='试验条件A', 合格标准='试验标准A')
        with self.assertRaisesRegex(app.UserError, '证据'):
            self.save('测试轮次', dict(full, 状态='通过'))
        first = self.save('测试轮次', dict(full, 状态='不通过', 测试结果与原因='未达到标准', 证据路径或链接='测试报告A.pdf'))['id']
        second = self.save('测试轮次', dict(full, 状态='通过', 测试结果与原因='达到标准', 证据路径或链接='测试报告B.pdf'))['id']
        reopened = app.Store(self.store.path).read()[1]['data']['测试轮次']
        self.assertEqual([r['_id'] for r in reopened], [first, second])
        self.assertEqual([r['状态'] for r in reopened], ['不通过', '通过'])
        self.assertEqual(reopened[0]['测试结果与原因'], '未达到标准')

    def test_sample_followup_atomic_idempotent_and_two_way_dates(self):
        cid = self.company()
        project = self.save('用胶项目', {'公司编号': cid})['id']
        values = {'项目编号': project, '完整型号': '联动测试胶', '状态': '已寄出', '数量': 1, '单位': '支', '寄出日期': '2026-10-18'}
        old = app.digest(self.store.path)
        with self.assertRaisesRegex(app.UserError, '回访'):
            self.save('样品测试', values)
        self.assertEqual(old, app.digest(self.store.path))
        values.update({'约定反馈日期': '2026-10-22', '日期依据': '本人计划'})
        sid = self.save('样品测试', values)['id']
        data = self.store.read()[1]['data']
        sample = data['样品测试'][0]
        tid = sample['任务编号']
        task = next(t for t in data['跟进待办'] if t['_id'] == tid)
        self.assertEqual(task['样品编号'], sid)
        self.assertEqual(task['项目编号'], project)
        self.assertEqual(task['计划日期'], '2026-10-22')
        history = sample['推进记录']
        self.save('样品测试', values, sid)
        self.assertEqual(len(self.store.read()[1]['data']['跟进待办']), 1)
        self.assertEqual(self.store.read()[1]['data']['样品测试'][0]['推进记录'], history)
        old = app.digest(self.store.path)
        with self.assertRaisesRegex(app.UserError, '延期原因'):
            self.save('样品测试', {'约定反馈日期': '2026-10-23'}, sid)
        self.assertEqual(old, app.digest(self.store.path))
        self.save('样品测试', {'约定反馈日期': '2026-10-23', '延期原因': '客户安排延后'}, sid)
        task = self.store.read()[1]['data']['跟进待办'][0]
        self.assertEqual(task['计划日期'], '2026-10-23')
        self.assertIn('2026-10-22', task['日期调整记录'])
        self.save('跟进待办', {'计划日期': '2026-10-24', '延期原因': '调整本人计划'}, tid)
        sample = self.store.read()[1]['data']['样品测试'][0]
        self.assertEqual(sample['约定反馈日期'], '2026-10-24')
        self.assertIn('2026-10-23', sample['推进记录'])
        self.assertEqual(len(self.store.read()[1]['data']['跟进待办']), 1)

    def test_sample_application_format_fields(self):
        cid = self.company()
        project = self.save('用胶项目', {'公司编号': cid})['id']
        values = {'项目编号': project, '状态': '已申请', '申请时间': '2026-05-22', '图号': 'DRAW-002', '物料名称': '黄胶', '完整型号': '', '规格': '100 mL/支', '数量': 2, '单位': '支', '业务员': '格式验收', '备注': '只在测试副本保存'}
        old = app.digest(self.store.path)
        with self.assertRaisesRegex(app.UserError, '有效日期'):
            self.save('样品测试', dict(values, 申请时间='2026-99-22'))
        self.assertEqual(old, app.digest(self.store.path))
        rid = self.save('样品测试', values)['id']
        data = app.Store(self.store.path).read()[1]['data']
        actual = next(x for x in data['样品测试'] if x['_id'] == rid)
        for key, value in values.items():
            self.assertEqual(actual[key], value, key)
        self.assertEqual(len(data['跟进待办']), 1)
        self.assertEqual(actual['任务编号'], data['跟进待办'][0]['_id'])

    def test_third_china_workday_and_sample_schedule(self):
        self.assertEqual(app.third_workday('2026-10-16'), '2026-10-21')
        self.assertEqual(app.third_workday('2026-10-18'), '2026-10-21')
        self.assertEqual(app.third_workday('2026-09-30'), '2026-10-10')
        self.assertEqual(app.third_workday('2026-02-13'), '2026-02-25')
        self.assertEqual(app.third_workday('2026-05-22'), '2026-05-27')
        self.assertEqual(app.third_workday('2026-12-31'), '')
        cid = self.company()
        project = self.save('用胶项目', {'公司编号': cid})['id']
        sid = self.save('样品测试', {'项目编号': project, '状态': '已申请', '完整型号': '排期验证', '签收日期': '2026-09-30'})['id']
        sample = self.store.read()[1]['data']['样品测试'][0]
        self.assertEqual(sample['计划测试日期'], '2026-10-10')
        self.assertEqual(sample['约定反馈日期'], '2026-10-10')
        task = self.store.read()[1]['data']['跟进待办'][0]
        self.assertEqual(task['计划日期'], '2026-10-10')
        self.assertEqual(task['日期依据'], '本人计划')
        second = self.save('样品测试', {'项目编号': project, '状态': '已申请', '完整型号': '人工排期', '签收日期': '2026-09-30', '计划测试日期': '2026-10-12', '约定反馈日期': '2026-10-14', '日期依据': '客户约定'})['id']
        row = next(r for r in self.store.read()[1]['data']['样品测试'] if r['_id'] == second)
        self.assertEqual(row['计划测试日期'], '2026-10-12')
        self.assertEqual(row['约定反馈日期'], '2026-10-14')

    def test_polished_visit_preserves_original(self):
        cid = self.company()
        values = {'公司编号': cid, '日期': '2026-09-21', '渠道': '面访', '客户原话及时间戳': '00:05 技术甲：原始重复重复。', '我方原话及时间戳': '00:10 我自己：我我介绍。', '客户整理稿及时间戳': '00:05 技术甲：原始重复。', '我方整理稿及时间戳': '00:10 我自己：我介绍。', '整理稿路径': '/local/整理稿.md'}
        rid = self.save('沟通记录', values)['id']
        reopened = app.Store(self.store.path).read()[1]['data']['沟通记录'][0]
        for key, value in values.items():
            self.assertEqual(reopened[key], value)

    def test_project_delete_cascade_preserves_sources_and_backup(self):
        cid = self.company()
        pid = self.save('用胶项目', {'公司编号': cid})['id']
        keep = self.save('用胶项目', {'公司编号': cid, '终端产品': '保留项目'})['id']
        sid = self.save('样品测试', {'项目编号': pid, '完整型号': '删除副本样品'})['id']
        self.save('测试轮次', {'项目编号': pid, '样品编号': sid, '日期': '2026-10-19', '状态': '待安排'})
        vid = self.save('沟通记录', {'公司编号': cid, '项目编号': pid, '日期': '2026-10-19', '客户原话及时间戳': '原文必须保留'})['id']
        kid = self.save('技术经验', {'项目编号': pid, '客户问题原话': '经验必须保留'})['id']
        old = self.store.read()[1]
        result = self.store.delete_sample({'sheet': '用胶项目', 'id': pid, 'version': old['version']})
        after = self.store.read()[1]['data']
        self.assertEqual([r['_id'] for r in after['用胶项目']], [keep])
        self.assertFalse(after['样品测试'])
        self.assertFalse(after['测试轮次'])
        self.assertFalse(after['跟进待办'])
        self.assertEqual(after['公司'][0]['_id'], cid)
        self.assertEqual(after['沟通记录'][0]['_id'], vid)
        self.assertEqual(after['沟通记录'][0]['客户原话及时间戳'], '原文必须保留')
        self.assertEqual(after['沟通记录'][0]['项目编号'], '')
        self.assertEqual(after['技术经验'][0]['_id'], kid)
        self.assertEqual(after['技术经验'][0]['项目编号'], '')
        backed = app.Store(Path(result['backup'])).read()[1]['data']
        self.assertEqual(len(backed['用胶项目']), 2)
        self.assertEqual(len(backed['测试轮次']), 1)
        with self.assertRaisesRegex(app.UserError, '刷新'):
            self.store.delete_sample({'sheet': '用胶项目', 'id': keep, 'version': old['version']})

    def test_business_row_updates_samples_once_and_rolls_back(self):
        cid = self.company('示例公司')
        self.save('联系人', {'公司编号': cid, '姓名': '陈采购', '职位': '采购'})
        pid = self.save('用胶项目', {'公司编号': cid, '终端产品': '示例设备', '用胶部位': '外壳密封'})['id']
        common = {'项目编号': pid, '数量': 1, '单位': '支', '寄出日期': '2026-10-18', '签收日期': '2026-10-18', '日期依据': '本人计划'}
        received = self.save('样品测试', dict(common, 完整型号='DEMO-A', 状态='已寄出', 约定反馈日期='2026-10-18', 问题与下一步='导入占位：已申请。未填写单号'))['id']
        kept = self.save('样品测试', dict(common, 完整型号='DEMO-B', 状态='已寄出', 约定反馈日期='2026-10-23', 问题与下一步='完整型号 DEMO-OLD。数量、快递单号未记录。'))['id']
        waiting = self.save('样品测试', {'项目编号': pid, '完整型号': 'DEMO-C', '状态': '已申请', '问题与下一步': '导入占位，未填写单号'})['id']
        other = self.save('用胶项目', {'公司编号': cid, '终端产品': '不要改'})['id']
        other_sample = self.save('样品测试', {'项目编号': other, '完整型号': '别动这款', '状态': '测试中', '数量': 1, '单位': '支', '寄出日期': '2026-10-01', '签收日期': '2026-10-02', '约定反馈日期': '2026-10-26', '日期依据': '客户约定'})['id']
        state = self.store.read()[1]
        old_digest = state['version']
        with self.assertRaisesRegex(app.UserError, '刷新'):
            self.store.save_project({'id': pid, 'version': 'stale', 'samples': [{'id': received, '状态': '已签收'}]})
        self.assertEqual(old_digest, app.digest(self.store.path))
        with self.assertRaisesRegex(app.UserError, '寄出日期'):
            self.store.save_project({'id': pid, 'version': old_digest, 'samples': [{'id': waiting, '状态': '已寄出'}]})
        self.assertEqual(old_digest, app.digest(self.store.path))
        self.store.save_project({'id': pid, 'version': old_digest, 'project': {'最新反馈': '两款已收样，开始安排测试', '客户问题与目标': '', '项目备注': '迁移记录：原客户问题与目标为样品跟踪器导入。'}, 'samples': [
            {'id': received, '状态': '已签收', '约定反馈日期': '2026-10-21', '日期依据': '本人计划', '延期原因': '按收样日重排测试', '问题与下一步': '', '迁移记录': '原状态已寄出；原下一步含未填写单号'},
            {'id': kept, '状态': '已签收', '问题与下一步': '', '备注': '待核：样品栏型号为 DEMO-B，历史备注写过 DEMO-OLD。', '迁移记录': '原下一步写过 DEMO-OLD'},
            {'id': waiting, '问题与下一步': '', '迁移记录': '原下一步为导入占位'},
        ]})
        again = self.store.read()[1]
        self.store.save_project({'id': pid, 'version': again['version'], 'samples': [
            {'id': received, '状态': '已签收'}, {'id': kept, '状态': '已签收'}, {'id': waiting, '状态': '已申请'},
        ]})
        data = self.store.read()[1]['data']
        samples = {r['_id']: r for r in data['样品测试']}
        self.assertEqual(samples[received]['状态'], '已签收')
        self.assertEqual(samples[received]['计划测试日期'], '2026-10-21')
        self.assertEqual(samples[received]['约定反馈日期'], '2026-10-21')
        self.assertNotIn('未填写单号', samples[received]['问题与下一步'])
        self.assertIn('迁移记录', samples[received]['推进记录'])
        self.assertEqual(samples[kept]['完整型号'], 'DEMO-B')
        self.assertEqual(samples[kept]['计划测试日期'], '2026-10-21')
        self.assertEqual(samples[kept]['约定反馈日期'], '2026-10-23')
        self.assertIn('DEMO-OLD', samples[kept]['推进记录'])
        self.assertIn('待核', samples[kept]['备注'])
        self.assertEqual(samples[waiting]['状态'], '已申请')
        self.assertEqual(samples[other_sample]['状态'], '测试中')
        tasks = [t for t in data['跟进待办'] if t.get('项目编号') == pid]
        self.assertEqual(len(tasks), 3)
        by_sample = {t['样品编号']: t for t in tasks}
        self.assertEqual(by_sample[received]['下一步具体动作'], 'DEMO-A · 安排测试')
        self.assertNotIn('确认签收', by_sample[received]['事项类型'])
        self.assertEqual(by_sample[kept]['下一步具体动作'], 'DEMO-B · 安排测试')
        self.assertEqual(by_sample[kept]['计划日期'], '2026-10-23')
        self.assertEqual(by_sample[waiting]['下一步具体动作'], 'DEMO-C · 跟进寄出')
        self.assertNotIn('通过', [samples[received]['状态'], samples[kept]['状态']])
        row = next(r for r in self.store.read()[1]['businessRows'] if r['id'] == pid)
        self.assertIn('已收样', row['sampleSummary'])
        self.assertIn('已申请', row['sampleSummary'])
        self.assertEqual(row['models'], ['DEMO-A', 'DEMO-B', 'DEMO-C'])
        self.assertNotIn('示例客户甲', ''.join(str(c or '') for c in row['cells']))
        book = load_workbook(self.store.path)
        self.assertEqual(book.sheetnames[0], '业务跟进表')
        headers = [c.value for c in book['业务跟进表'][2]]
        self.assertEqual(len(headers), 24)
        self.assertEqual(headers[0], '编号')
        text = '\n'.join(str(c.value or '') for row in book['业务跟进表'].iter_rows(min_row=3, max_col=24) for c in row)
        self.assertIn('DEMO-A', text)
        self.assertNotIn('示例客户甲', text)
        self.assertNotIn('测试已完成', text)
